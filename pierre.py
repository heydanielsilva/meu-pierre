"""Meu Pierre — núcleo: sincronização e análises de finanças pessoais usadas pelo painel.

Lê suas contas pelo Open Finance via Meu Pluggy (gratuito para uso pessoal),
guarda um cache no banco (SQLite no computador, Postgres na Vercel) e calcula
os números que o painel mostra: saldos, gastos, orçamentos, assinaturas e investimentos.

SOMENTE LEITURA: nada aqui movimenta dinheiro.
"""
from __future__ import annotations

import calendar
import json
import os
import re
import statistics
import uuid
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from pluggy_client import PluggyClient
from store import NEUTRAS, VERSAO_TRANSACOES, Store, _sem_acento, agora_br, hoje_br

PASTA = Path(__file__).resolve().parent


# ------------------------------------------------------------------ config
def _carregar_env() -> None:
    arq = PASTA / ".env"
    if not arq.exists():
        return
    for linha in arq.read_text(encoding="utf-8").splitlines():
        linha = linha.strip()
        if not linha or linha.startswith("#") or "=" not in linha:
            continue
        k, v = linha.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


_carregar_env()
DEMO = os.environ.get("MEU_PIERRE_DEMO", "0") == "1"
NUVEM = bool(os.environ.get("VERCEL"))  # a Vercel define VERCEL=1
SALDO_MINIMO = float(os.environ.get("SALDO_MINIMO") or "500")


def _itens() -> dict[str, str]:
    """PLUGGY_ITEMS=itemId:Nubank,itemId2:Itaú  ->  {itemId: 'Nubank', ...}"""
    out = {}
    for parte in os.environ.get("PLUGGY_ITEMS", "").split(","):
        parte = parte.strip()
        if not parte:
            continue
        item_id, _, nome = parte.partition(":")
        out[item_id.strip()] = nome.strip()
    return out


def _db_url() -> str:
    """DATABASE_URL ou POSTGRES_URL, inclusive com o prefixo que a Vercel adiciona ao conectar o Neon
    (ex.: DATABASE_URL_DATABASE_URL, STORAGE_POSTGRES_URL)."""
    postgres = ("postgres://", "postgresql://")
    for nome in ("DATABASE_URL", "POSTGRES_URL"):
        if os.environ.get(nome, "").startswith(postgres):
            return os.environ[nome]
    for nome, valor in sorted(os.environ.items()):
        if nome.endswith(("DATABASE_URL", "POSTGRES_URL")) and valor.startswith(postgres):
            return valor
    return ""


DB_URL = _db_url()
if DB_URL and not DEMO:
    store = Store(DB_URL)
elif NUVEM and not DEMO:
    raise RuntimeError("Na Vercel é preciso um banco Postgres: conecte um banco Neon em Storage e faça Redeploy.")
else:
    db_nome = "/tmp/meu_pierre_demo.db" if (DEMO and NUVEM) else PASTA / ("meu_pierre_demo.db" if DEMO else "meu_pierre.db")
    store = Store(os.environ.get("MEU_PIERRE_DB", db_nome))
if DEMO:
    store.popular_demo()


# ----------------------------------------------------------------- helpers
def brl(v: float | None) -> str:
    if v is None:
        return "—"
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"-R$ {s}" if v < 0 else f"R$ {s}"


def _mes(mes: str = "") -> tuple[str, str, str]:
    """'' (mês atual) ou 'AAAA-MM'  ->  ('AAAA-MM', primeiro dia, último dia)."""
    hoje = hoje_br()
    ano, num = hoje.year, hoje.month
    if mes:
        if not re.fullmatch(r"\d{4}-\d{1,2}", mes.strip()):
            raise ValueError(f"Mês inválido '{mes}'. Use AAAA-MM, ex.: 2026-10.")
        ano, num = map(int, mes.strip().split("-"))
    ult = calendar.monthrange(ano, num)[1]
    return f"{ano}-{num:02d}", f"{ano}-{num:02d}-01", f"{ano}-{num:02d}-{ult}"


def _mes_anterior(rotulo: str) -> str:
    return (date.fromisoformat(f"{rotulo}-01") - timedelta(days=1)).strftime("%Y-%m")


def garantir_dados() -> str:
    """Sincroniza sozinho se nunca sincronizou ou se faz mais de 24h. Devolve um aviso se falhar."""
    if DEMO:
        return ""
    ult = store.get_meta("ultima_sync")
    limite = timedelta(hours=30 if NUVEM else 24)  # na nuvem o cron diário cuida disso
    regra_nova = bool(ult) and store.get_meta("versao_transacoes") != VERSAO_TRANSACOES
    if ult and agora_br() - datetime.fromisoformat(ult) < limite and not regra_nova:
        return ""
    try:
        _sync(90 if NUVEM and not ult else 365 if not ult else 30)
        return ""
    except Exception as e:  # noqa: BLE001
        return f"Não consegui atualizar os dados agora: {e}. Mostrando o último cache."


def _sync(dias: int) -> dict:
    try:
        itens = _itens()
        if not itens:
            raise RuntimeError("Nenhum item configurado em PLUGGY_ITEMS.")
        cli = PluggyClient(os.environ.get("PLUGGY_CLIENT_ID", ""), os.environ.get("PLUGGY_CLIENT_SECRET", ""))
        return store.sincronizar(cli, itens, dias)
    except Exception as e:
        # Guarda o motivo para o painel mostrar, inclusive quando a atualização foi automática.
        store.registrar_resultado({"em": agora_br().isoformat(), "dias": dias, "contas": 0, "transacoes": 0,
                                   "investimentos": 0, "faturas": 0, "erros": [str(e)], "status": {}})
        raise


def sincronizar(dias: int = 30) -> dict:
    """Busca dados novos no Open Finance (Meu Pluggy). O Meu Pluggy atualiza os bancos 1x por dia."""
    if DEMO:
        return {"contas": 0, "transacoes": 0, "investimentos": 0, "faturas": 0,
                "erros": ["Modo demonstração: os dados são fictícios e não há o que sincronizar."]}
    return _sync(max(1, min(dias, 365)))


def _fluxo(txs: list[dict]) -> tuple[float, float]:
    """Receitas e gastos reais. Estornos abatem dos gastos em vez de contar como receita."""
    validas = [t for t in txs if not t["neutra"]]
    rec = sum(t["valor"] for t in validas if t["direcao"] == "entrada" and t["categoria"] != "Estornos")
    estornos = sum(t["valor"] for t in validas if t["direcao"] == "entrada" and t["categoria"] == "Estornos")
    desp = sum(t["valor"] for t in validas if t["direcao"] == "saida") - estornos
    return rec, desp


def _por_categoria(txs: list[dict]) -> dict[str, float]:
    tot: dict[str, float] = defaultdict(float)
    for t in txs:
        if t["neutra"]:
            continue
        if t["direcao"] == "saida":
            tot[t["categoria"]] += t["valor"]
        elif t["categoria"] == "Estornos":
            tot["Estornos"] -= t["valor"]  # aparece negativo e abate o total
    return {k: v for k, v in sorted(tot.items(), key=lambda kv: -kv[1]) if abs(v) > 0.005}


def _chave_recorrente(t: dict) -> str:
    base = t.get("estabelecimento") or t["descricao"]
    base = _sem_acento(base)
    base = re.sub(r"[\d*#.\-_/]+", " ", base)
    return " ".join(base.split()[:3])


# ================================================================ ANÁLISES
def meses_disponiveis() -> list[str]:
    """Meses com transações no cache, do mais recente para o mais antigo (inclui sempre o atual)."""
    r = store.con.execute("SELECT MIN(data) de FROM transactions").fetchone()
    atual = hoje_br().replace(day=1)
    m = date.fromisoformat(r["de"][:7] + "-01") if r and r["de"] else atual
    out = []
    while m <= atual:
        out.append(m.strftime("%Y-%m"))
        m = (m + timedelta(days=32)).replace(day=1)
    return out[::-1]


def status_conexoes() -> list[dict]:
    """Cada banco conectado: última atualização, erro e quando a autorização expira."""
    hoje = hoje_br()
    out = []
    for r in store.con.execute("SELECT * FROM items").fetchall():
        exp = r["consentimento_expira"]
        out.append({"banco": r["banco"] or r["id"], "status": r["status"],
                    "atualizado": r["ultima_atualizacao"], "expira": exp[:10] if exp else None,
                    "dias_para_expirar": (date.fromisoformat(exp[:10]) - hoje).days if exp else None,
                    "erro": r["erro"]})
    return out


def contas() -> list[dict]:
    """Contas correntes, poupanças e cartões com saldo, limite e vencimento."""
    out = []
    for c in store.con.execute("SELECT * FROM accounts ORDER BY banco, tipo").fetchall():
        cartao = c["subtipo"] == "CREDIT_CARD"
        out.append({"id": c["id"], "banco": c["banco"], "nome": c["nome"], "cartao": cartao, "tipo": c["tipo"],
                    "final": str(c["numero"] or "")[-4:], "saldo": c["saldo"], "limite": c["limite"],
                    "limite_disponivel": c["limite_disponivel"], "vencimento": c["vencimento_fatura"],
                    "fechamento": c["fechamento_fatura"],
                    "uso_limite": (1 - (c["limite_disponivel"] or 0) / c["limite"]) if cartao and c["limite"] else None,
                    **(_faturas_do_cartao(c) if cartao else {})})
    return out


def _mais_meses(d: date, n: int) -> date:
    m = d.month - 1 + n
    a, m = d.year + m // 12, m % 12 + 1
    return date(a, m, min(d.day, calendar.monthrange(a, m)[1]))


def _faturas_do_cartao(c, diagnostico: bool = False) -> dict:
    """Fatura aberta, fatura fechada a pagar e as próximas (parcelas já lançadas) de um cartão.

    O saldo que a Pluggy manda para o cartão é o total devido (inclui parcelas futuras), e /bills só
    traz faturas fechadas. Por isso a fatura aberta é calculada pelas compras:
    - se o banco informa em que fatura cada compra vai cair (billForecastDate, "AAAA-MM"): a aberta é
      a primeira fatura prevista a partir deste mês e as seguintes são as próximas (as parcelas futuras
      já vêm lançadas como pendentes, cada uma com a sua fatura prevista);
    - se o banco informa a fatura de cada compra (billId): soma as faturas que ainda não fecharam,
      isto é, que não estão em /bills e têm compra depois do último fechamento (as parcelas de
      compras antigas vêm junto, porque têm o mesmo billId); compras sem billId entram pela data;
    - se não informa: soma as compras do ciclo pela data em que entram na fatura. Parcelas que
      vêm com a data da compra original são empurradas para o mês da parcela.
    """
    hoje = hoje_br()
    venc = date.fromisoformat(c["vencimento_fatura"]) if c["vencimento_fatura"] else None
    fech_banco = date.fromisoformat(c["fechamento_fatura"]) if c["fechamento_fatura"] else None
    fech = fech_banco or (venc - timedelta(days=7) if venc else None)
    if not fech:
        fech = hoje.replace(day=calendar.monthrange(hoje.year, hoje.month)[1])
    folga = max(1, (venc - fech).days) if venc and venc > fech else 7
    while fech < hoje:                                   # próximo fechamento a partir de hoje
        fech = _mais_meses(fech, 1)
    while _mais_meses(fech, -1) >= hoje:
        fech = _mais_meses(fech, -1)
    ini = _mais_meses(fech, -1)                          # ciclo aberto: (ini, fech]
    ini_ant = _mais_meses(ini, -1)                       # ciclo anterior: (ini_ant, ini]

    txs = [dict(t) for t in store.con.execute(
        "SELECT * FROM transactions WHERE account_id=? AND data > ?", (c["id"], (hoje - timedelta(days=400)).isoformat())).fetchall()]
    fechadas_banco = {b["id"]: dict(b) for b in store.con.execute("SELECT * FROM bills WHERE account_id=?", (c["id"],)).fetchall()}

    def valor(t) -> float:
        if t["direcao"] == "saida":
            return t["valor"]
        return -t["valor"] if "estorno" in _sem_acento(t["descricao"] or "") else 0.0   # pagamento da fatura não conta

    # data em que a compra entra na fatura: parcela que veio com a data da compra original vai para o mês dela
    datas_por_compra: dict[tuple, set] = defaultdict(set)
    for t in txs:
        if (t.get("parcelas_total") or 0) > 1:
            datas_por_compra[(_chave_recorrente(t), round(t["valor"], 2), t["parcelas_total"])].add((t["parcela_atual"], t["data"]))

    def data_fatura(t) -> str:
        p = t.get("parcela_atual") or 0
        if (t.get("parcelas_total") or 0) <= 1 or p <= 1:
            return t["data"]
        irmas = datas_por_compra[(_chave_recorrente(t), round(t["valor"], 2), t["parcelas_total"])]
        mesma_data = len({d for _, d in irmas}) == 1 and len(irmas) > 1
        compra = t.get("data_compra")
        if mesma_data or (compra and abs((date.fromisoformat(t["data"]) - date.fromisoformat(compra)).days) <= 3):
            return _mais_meses(date.fromisoformat(t["data"]), p - 1).isoformat()
        return t["data"]

    def no_ciclo(t, de: date, ate: date) -> bool:
        return de.isoformat() < data_fatura(t) <= ate.isoformat()

    grupos: dict[str, list[dict]] = defaultdict(list)
    for t in txs:
        if t.get("fatura_id"):
            grupos[t["fatura_id"]].append(t)
    com_fatura = bool(grupos)
    # fatura prevista pelo banco para o que ainda não foi faturado (sem billId)
    previstas: dict[str, float] = defaultdict(float)
    for t in txs:
        if not t.get("fatura_id") and t.get("fatura_prevista"):
            previstas[t["fatura_prevista"][:7]] += valor(t)
    # a aberta é a que vence neste ciclo; se o banco usar outro mês de referência, a primeira a partir de hoje
    mes_venc = (fech + timedelta(days=folga)).strftime("%Y-%m")
    mes_aberta = mes_venc if mes_venc in previstas else min((m for m in previstas if m >= hoje.strftime("%Y-%m")), default=None)
    if mes_aberta:
        aberta = previstas[mes_aberta] + sum(valor(t) for t in txs if not t.get("fatura_id") and not t.get("fatura_prevista")
                                             and no_ciclo(t, ini, fech))
        abertas_ids = set()
    elif com_fatura:
        abertas_ids = {g for g, ts in grupos.items() if g not in fechadas_banco and max(x["data"] for x in ts) > ini.isoformat()}
        aberta = sum(valor(t) for t in txs if t.get("fatura_id") in abertas_ids or (not t.get("fatura_id") and no_ciclo(t, ini, fech)))
    else:
        abertas_ids = set()
        aberta = sum(valor(t) for t in txs if no_ciclo(t, ini, fech))

    # fatura já fechada que ainda vence (a do ciclo anterior)
    fechada = None
    venc_ant = ini + timedelta(days=folga)
    if venc_ant >= hoje:
        b = store.con.execute("SELECT * FROM bills WHERE account_id=? AND vencimento >= ? ORDER BY vencimento LIMIT 1",
                              (c["id"], hoje.isoformat())).fetchone()
        if b:
            fechada = {"valor": b["valor"], "vence": b["vencimento"], "minimo": b["pagamento_minimo"]}
        else:
            anteriores = [m for m in previstas if mes_aberta and m < mes_aberta and m >= _mais_meses(hoje, -1).strftime("%Y-%m")]
            if anteriores:
                v = previstas[max(anteriores)]
            elif com_fatura:
                ant_ids = {g for g, ts in grupos.items() if g not in abertas_ids and g not in fechadas_banco
                           and ini_ant.isoformat() < max(x["data"] for x in ts) <= ini.isoformat()}
                v = sum(valor(t) for t in txs if t.get("fatura_id") in ant_ids or (not t.get("fatura_id") and no_ciclo(t, ini_ant, ini)))
            else:
                v = sum(valor(t) for t in txs if no_ciclo(t, ini_ant, ini))
            if v > 0:
                fechada = {"valor": round(v, 2), "vence": venc_ant.isoformat(), "minimo": None}

    # próximas faturas: só o que já está comprometido em parcelas deste cartão
    parcelas = [p for p in parcelas_futuras() if p.get("conta_id") == c["id"]]
    if mes_aberta:
        base = date.fromisoformat(mes_aberta + "-01")
        proximas = [{"vence": (_mais_meses(fech, k) + timedelta(days=folga)).isoformat(),
                     "valor": round(previstas.get(_mais_meses(base, k).strftime("%Y-%m"), 0.0), 2)} for k in range(1, 7)]
    else:
        proximas = [{"vence": (_mais_meses(fech, k) + timedelta(days=folga)).isoformat(),
                     "valor": sum(p["valor"] for p in parcelas if p["restantes"] >= k)} for k in range(1, 7)]
    out = {"fatura_aberta": round(aberta, 2), "fecha_em": fech.isoformat(), "vence_em": (fech + timedelta(days=folga)).isoformat(),
           "fatura_fechada": fechada, "proximas_faturas": proximas, "total_devido": c["saldo"]}
    if diagnostico:  # só números e datas, sem descrições
        out["diagnostico"] = {
            "fechamento_banco": c["fechamento_fatura"], "vencimento_banco": c["vencimento_fatura"],
            "ciclo_aberto": [ini.isoformat(), fech.isoformat()], "ciclo_anterior": [ini_ant.isoformat(), ini.isoformat()],
            "metodo": "fatura prevista" if mes_aberta else "billId" if com_fatura else "datas",
            "fatura_aberta_prevista": mes_aberta, "com_fatura_prevista": sum(1 for t in txs if t.get("fatura_prevista")),
            "data_mais_nova": max((t["data"] for t in txs), default=None),
            "soma_por_fatura_prevista": {m: round(v, 2) for m, v in sorted(previstas.items())}, "transacoes_400_dias": len(txs),
            "com_billId": sum(1 for t in txs if t.get("fatura_id")), "com_data_compra": sum(1 for t in txs if t.get("data_compra")),
            "parceladas": sum(1 for t in txs if (t.get("parcelas_total") or 0) > 1),
            "faturas_fechadas_do_banco": [{"vence": b["vencimento"], "valor": b["valor"]} for b in
                                          sorted(fechadas_banco.values(), key=lambda b: b["vencimento"] or "", reverse=True)[:4]],
            "grupos_por_billId": sorted([{"id": g[-6:], "qtd": len(ts), "de": min(x["data"] for x in ts), "ate": max(x["data"] for x in ts),
                                          "soma": round(sum(valor(x) for x in ts), 2), "fechada_no_banco": g in fechadas_banco,
                                          "conta_como_aberta": g in abertas_ids} for g, ts in grupos.items()],
                                        key=lambda x: x["ate"], reverse=True)[:6],
            "soma_por_data_ciclo_aberto": round(sum(valor(t) for t in txs if no_ciclo(t, ini, fech)), 2),
            "soma_por_data_ciclo_anterior": round(sum(valor(t) for t in txs if no_ciclo(t, ini_ant, ini)), 2),
            "parcelas_recentes": [{"data": t["data"], "data_compra": t.get("data_compra"), "parcela": f"{t['parcela_atual']}/{t['parcelas_total']}",
                                   "valor": t["valor"], "entra_na_fatura_em": data_fatura(t), "billId": (t.get("fatura_id") or "")[-6:]}
                                  for t in sorted(txs, key=lambda t: t["data"], reverse=True) if (t.get("parcelas_total") or 0) > 1][:10],
        }
    return out


def diagnostico_sync() -> dict:
    """Versão das transações gravadas e o resultado da última sincronização."""
    try:
        ultimo = json.loads(store.get_meta("ultimo_resultado") or "{}")
    except ValueError:
        ultimo = {}
    return {"versao_gravada": store.get_meta("versao_transacoes"), "versao_atual": VERSAO_TRANSACOES,
            "ultima_sync": store.get_meta("ultima_sync"), "ultimo_resultado": ultimo}


def diagnostico_cartoes() -> list[dict]:
    """Resumo do cálculo das faturas de cada cartão, para conferir com o app do banco (sem descrições)."""
    out = []
    for c in store.con.execute("SELECT * FROM accounts WHERE subtipo='CREDIT_CARD'").fetchall():
        f = _faturas_do_cartao(c, diagnostico=True)
        out.append({"cartao": f"{c['banco']} final {str(c['numero'] or '')[-4:]}", "total_devido_pluggy": c["saldo"],
                    "limite": c["limite"], "limite_disponivel": c["limite_disponivel"],
                    "fatura_aberta_calculada": f["fatura_aberta"], "fatura_fechada": f["fatura_fechada"], **f["diagnostico"]})
    return out


def patrimonio() -> dict:
    lst = contas()
    saldo = sum(c["saldo"] or 0 for c in lst if c["tipo"] == "BANK")
    faturas = sum(c["saldo"] or 0 for c in lst if c["cartao"])
    invs = store.con.execute("SELECT SUM(saldo) s FROM investments").fetchone()["s"] or 0
    return {"saldo": saldo, "investimentos": invs, "faturas": faturas, "liquido": saldo + invs - faturas}


def resumo_mes(mes: str = "") -> dict:
    """Receitas, gastos, sobra, taxa de poupança, categorias e estabelecimentos, comparando com o mês anterior."""
    rotulo, de, ate = _mes(mes)
    txs = store.transacoes(de, ate)
    rec, desp = _fluxo(txs)
    _, de_a, ate_a = _mes(_mes_anterior(rotulo))
    txs_a = store.transacoes(de_a, ate_a)
    rec_a, desp_a = _fluxo(txs_a)
    cats, cats_a = _por_categoria(txs), _por_categoria(txs_a)
    estab: dict[str, float] = defaultdict(float)
    for t in txs:
        if t["direcao"] == "saida" and not t["neutra"]:
            estab[t.get("estabelecimento") or t["descricao"]] += t["valor"]
    hoje = hoje_br()
    atual = rotulo == hoje.strftime("%Y-%m")
    return {
        "mes": rotulo, "atual": atual, "dia": hoje.day if atual else None,
        "dias_mes": calendar.monthrange(int(rotulo[:4]), int(rotulo[5:]))[1],
        "receitas": rec, "gastos": desp, "sobra": rec - desp,
        "poupanca": (rec - desp) / rec if rec else None,
        "receitas_anterior": rec_a, "gastos_anterior": desp_a,
        "projecao": _projecao_mes(txs) if atual else None,
        "categorias": [{"categoria": c, "valor": v, "anterior": cats_a.get(c, 0.0)} for c, v in cats.items()],
        "estabelecimentos": [{"nome": e, "valor": v} for e, v in sorted(estab.items(), key=lambda kv: -kv[1])[:10]],
        "transacoes": txs,
    }


def evolucao_mensal(meses: int = 12) -> list[dict]:
    """Receitas, gastos e sobra mês a mês nos últimos N meses (o último é o atual, parcial)."""
    m = hoje_br().replace(day=1)
    lista = []
    for _ in range(meses):
        lista.append(m.strftime("%Y-%m"))
        m = (m - timedelta(days=1)).replace(day=1)
    out = []
    for rot in reversed(lista):
        _, de, ate = _mes(rot)
        rec, desp = _fluxo(store.transacoes(de, ate))
        out.append({"mes": rot, "receitas": rec, "gastos": desp, "sobra": rec - desp})
    return out


CONTAS_FIXAS = {"Aluguel", "Contas da casa", "Internet e telefone", "Plano de saúde", "Seguros",
                "Educação", "Moradia", "Academia", "Previdência", "Empréstimos e financiamentos"}


def _recorrentes() -> list[dict]:
    """Cobranças que se repetem todo mês: valor fixo (assinatura) ou conta fixa de valor variável."""
    hoje = hoje_br()
    txs = store.transacoes((hoje - timedelta(days=185)).isoformat(), hoje.isoformat())
    grupos: dict[str, list[dict]] = defaultdict(list)
    for t in txs:
        if t["direcao"] == "saida" and not t["neutra"] and not t.get("parcelas_total"):
            grupos[_chave_recorrente(t)].append(t)
    achados = []
    for lst in grupos.values():
        meses = {t["data"][:7] for t in lst}
        ult = max(lst, key=lambda t: t["data"])
        d = date.fromisoformat(ult["data"])
        if len(meses) < 3 or len(lst) > len(meses) * 1.5 or (hoje - d).days > 45:
            continue  # 3+ meses, ~1x por mês e ainda ativa
        valores = [t["valor"] for t in lst]
        media = statistics.mean(valores)
        cv = statistics.pstdev(valores) / media if media else 0
        if ult["categoria"] in CONTAS_FIXAS and cv <= 0.40:
            tipo = "conta fixa"
        elif cv <= 0.05:
            tipo = "assinatura"
        else:
            continue
        prox = (d.replace(day=1) + timedelta(days=32)).replace(day=min(d.day, 28))
        achados.append({"nome": ult.get("estabelecimento") or ult["descricao"], "categoria": ult["categoria"],
                        "media": media, "ultima": d, "proxima": prox, "tipo": tipo,
                        "chave": _chave_recorrente(ult)})
    return sorted(achados, key=lambda a: -a["media"])


def recorrentes() -> list[dict]:
    """Assinaturas e contas fixas dos últimos 6 meses, prontas para o painel."""
    return [{k: (v.isoformat() if isinstance(v, date) else v) for k, v in r.items() if k != "chave"}
            for r in _recorrentes()]


def _projecao_mes(txs_mes: list[dict]) -> float:
    """Gasto atual + ritmo diário dos gastos variáveis + recorrentes que ainda vão cair no mês."""
    hoje = hoje_br()
    dias_mes = calendar.monthrange(hoje.year, hoje.month)[1]
    rec = _recorrentes()
    chaves = {r["chave"] for r in rec}
    gastos = [t for t in txs_mes if t["direcao"] == "saida" and not t["neutra"]]
    atual = sum(t["valor"] for t in gastos)
    variavel = sum(t["valor"] for t in gastos if _chave_recorrente(t) not in chaves and not t.get("parcelas_total"))
    ja_cobradas = {_chave_recorrente(t) for t in gastos}
    pendentes = sum(r["media"] for r in rec if r["chave"] not in ja_cobradas
                    and r["proxima"].month == hoje.month and r["proxima"].year == hoje.year)
    return atual + variavel / hoje.day * (dias_mes - hoje.day) + pendentes


def _somar_meses(d: date, n: int) -> str:
    m = d.month - 1 + n
    return f"{d.year + m // 12}-{m % 12 + 1:02d}"


def parcelas_futuras() -> list[dict]:
    """Compras parceladas em andamento: em que parcela estão, quanto falta e quando terminam."""
    rows = store.con.execute(
        """SELECT t.*, a.banco, a.subtipo FROM transactions t LEFT JOIN accounts a ON a.id = t.account_id
           WHERE t.parcelas_total > 1 AND t.data <= ? ORDER BY t.data DESC""", (hoje_br().isoformat(),)).fetchall()
    ultimas: dict = {}
    for r in rows:
        k = _chave_recorrente(dict(r))
        if k not in ultimas or r["parcela_atual"] > ultimas[k]["parcela_atual"]:
            ultimas[k] = r
    out = []
    for r in ultimas.values():
        rest = r["parcelas_total"] - r["parcela_atual"]
        if rest > 0:
            d = date.fromisoformat(r["data"][:10])
            out.append({"conta_id": r["account_id"], "descricao": r["descricao"], "estabelecimento": r["estabelecimento"], "banco": r["banco"],
                        "cartao": r["subtipo"] == "CREDIT_CARD", "parcela": r["parcela_atual"],
                        "total_parcelas": r["parcelas_total"], "valor": r["valor"], "restantes": rest,
                        "falta": rest * r["valor"], "total_compra": r["parcelas_total"] * r["valor"],
                        "ultima": r["data"][:10], "proxima": _somar_meses(d, 1), "termina": _somar_meses(d, rest)})
    return sorted(out, key=lambda p: -p["falta"])


def parcelas_por_mes(parcelas: list[dict], meses: int = 12) -> list[dict]:
    """Quanto das compras parceladas cai em cada um dos próximos meses."""
    hoje = hoje_br()
    soma: dict[str, float] = {_somar_meses(hoje, i): 0.0 for i in range(1, meses + 1)}
    for p in parcelas:
        d = date.fromisoformat(p["ultima"])
        for i in range(1, p["restantes"] + 1):
            m = _somar_meses(d, i)
            if m in soma:
                soma[m] += p["valor"]
    return [{"mes": m, "valor": v} for m, v in soma.items()]


# ================================================================ PREVISÃO
def adicionar_previsto(descricao: str, valor: float, tipo: str, mes: str, repetir: str, vezes: int | None) -> None:
    """Gasto ou entrada que você sabe que vai acontecer e o banco ainda não mostra."""
    if not descricao.strip() or valor <= 0 or not re.fullmatch(r"\d{4}-\d{2}", mes or ""):
        raise ValueError("Informe descrição, valor e mês.")
    tipo = "entrada" if tipo == "entrada" else "saida"
    repetir = repetir if repetir in ("unica", "mensal", "parcelado") else "unica"
    vezes = max(1, min(int(vezes), 120)) if vezes else None
    if repetir == "parcelado" and not vezes:
        raise ValueError("Informe em quantas parcelas.")
    store.con.execute("INSERT OR REPLACE INTO previstos VALUES (?,?,?,?,?,?,?,?)",
                      (uuid.uuid4().hex[:12], descricao.strip()[:80], round(valor, 2), tipo, mes, repetir,
                       vezes if repetir != "unica" else None, agora_br().isoformat()))
    store.con.commit()


def remover_previsto(id_: str) -> int:
    n = store.con.execute("DELETE FROM previstos WHERE id=?", (id_,)).rowcount
    store.con.commit()
    return n


def _valor_previsto_no_mes(p: dict, mes: str) -> float:
    """Quanto um lançamento manual pesa no mês (parcelado: valor total dividido nas parcelas)."""
    if mes < p["mes"]:
        return 0.0
    k = (int(mes[:4]) * 12 + int(mes[5:])) - (int(p["mes"][:4]) * 12 + int(p["mes"][5:]))
    if p["repetir"] == "unica":
        return p["valor"] if k == 0 else 0.0
    if p["repetir"] == "parcelado":
        return p["valor"] / p["vezes"] if k < p["vezes"] else 0.0
    return p["valor"] if not p["vezes"] or k < p["vezes"] else 0.0


def previsao(meses: int = 6) -> dict:
    """Este mês (projetado) e os próximos: contas fixas e assinaturas, parcelas já assumidas, gasto típico
    das outras categorias (mediana dos últimos 3 meses) e o que você adicionou à mão."""
    hoje = hoje_br()
    rec = _recorrentes()
    chaves = {r["chave"] for r in rec}
    fixos = sum(r["media"] for r in rec)

    # base: últimos 3 meses completos que têm transações
    por_cat: dict[str, list[float]] = defaultdict(list)
    receitas_base, base = [], []
    for k in (3, 2, 1):
        m = _somar_meses(hoje, -k)
        _, de, ate = _mes(m)
        txs = store.transacoes(de, ate)
        if not txs:
            continue
        base.append(m)
        receitas_base.append(_fluxo(txs)[0])
        cat: dict[str, float] = defaultdict(float)
        for t in txs:
            if t["neutra"] or t["direcao"] != "saida" or (t.get("parcelas_total") or 0) > 1 or _chave_recorrente(t) in chaves:
                continue
            cat[t["categoria"]] += t["valor"]
        for c, v in cat.items():
            por_cat[c].append(v)
    n = len(base)
    tipico = {c: statistics.median(vs + [0.0] * (n - len(vs))) for c, vs in por_cat.items()}
    tipico = {c: v for c, v in sorted(tipico.items(), key=lambda kv: -kv[1]) if v >= 1}
    variavel = sum(tipico.values())
    receita_tipica = statistics.median(receitas_base) if receitas_base else 0.0

    previstos = [dict(p) for p in store.con.execute("SELECT * FROM previstos ORDER BY mes, criado_em").fetchall()]
    parc = {m["mes"]: m["valor"] for m in parcelas_por_mes(parcelas_futuras(), meses)}

    out, acumulado = [], 0.0
    rotulo, de, ate = _mes("")
    txs_mes = store.transacoes(de, ate)
    for k in range(0, meses + 1):
        m = _somar_meses(hoje, k)
        manual_s = sum(_valor_previsto_no_mes(p, m) for p in previstos if p["tipo"] == "saida")
        manual_e = sum(_valor_previsto_no_mes(p, m) for p in previstos if p["tipo"] == "entrada")
        if k == 0:   # este mês: o que já foi gasto + o que ainda deve cair até o fim
            rec_ja, gasto_ja = _fluxo(txs_mes)
            proj = _projecao_mes(txs_mes)
            item = {"mes": m, "atual": True, "ja_gasto": round(gasto_ja, 2), "restante": round(max(0.0, proj - gasto_ja), 2),
                    "fixos": 0.0, "parcelas": 0.0, "variavel": 0.0,
                    "receitas": round(max(rec_ja, receita_tipica) + manual_e, 2), "receitas_ja": round(rec_ja, 2)}
            gastos = proj + manual_s
        else:
            p_m = parc.get(m, 0.0)
            item = {"mes": m, "atual": False, "ja_gasto": 0.0, "restante": 0.0, "fixos": round(fixos, 2),
                    "parcelas": round(p_m, 2), "variavel": round(variavel, 2), "receitas": round(receita_tipica + manual_e, 2)}
            gastos = fixos + p_m + variavel + manual_s
        item.update({"manuais": round(manual_s, 2), "manuais_entrada": round(manual_e, 2), "gastos": round(gastos, 2),
                     "sobra": round(item["receitas"] - gastos, 2)})
        acumulado += item["sobra"]
        item["acumulado"] = round(acumulado, 2)
        out.append(item)
    return {"meses": out, "base_meses": base, "fixos": round(fixos, 2), "receita_tipica": round(receita_tipica, 2),
            "variavel": round(variavel, 2), "variavel_categorias": [{"categoria": c, "valor": round(v, 2)} for c, v in tipico.items()],
            "lancamentos": [{k: p[k] for k in ("id", "descricao", "valor", "tipo", "mes", "repetir", "vezes")} for p in previstos]}


def faturas() -> list[dict]:
    """Faturas anteriores de cada cartão."""
    out = []
    for c in store.con.execute("SELECT * FROM accounts WHERE subtipo='CREDIT_CARD'").fetchall():
        for b in store.con.execute("SELECT * FROM bills WHERE account_id=? ORDER BY vencimento DESC LIMIT 12", (c["id"],)):
            out.append({"conta_id": c["id"], "cartao": f"{c['banco']} · {c['nome']}", "vencimento": b["vencimento"],
                        "valor": b["valor"], "minimo": b["pagamento_minimo"]})
    return out


CLASSES_INVEST = {"FIXED_INCOME": "Renda fixa", "MUTUAL_FUND": "Fundos", "EQUITY": "Ações", "ETF": "ETFs",
                  "SECURITY": "Previdência", "COE": "COE", "OTHER": "Outros"}
SUBTIPOS_INVEST = {"TREASURY": "Tesouro Direto", "REAL_ESTATE_FUND": "Fundos imobiliários", "CRYPTOCURRENCY": "Cripto"}


def _taxa(v, limite: float):
    """A Pluggy pode mandar a rentabilidade como fração (0,0095) ou em % (0,95); devolve sempre fração."""
    if v is None:
        return None
    return v / 100 if abs(v) > limite else v


def investimentos() -> list[dict]:
    """Carteira: saldo por aplicação, classe (renda fixa, ações…) e rentabilidade no mês e em 12 meses."""
    out = []
    for r in store.con.execute("SELECT * FROM investments ORDER BY saldo DESC").fetchall():
        classe = SUBTIPOS_INVEST.get(r["subtipo"] or "") or CLASSES_INVEST.get(r["tipo"] or "", "Outros")
        out.append({"id": r["id"], "nome": r["nome"], "banco": r["banco"], "tipo": r["tipo"], "subtipo": r["subtipo"], "classe": classe,
                    "saldo": r["saldo"], "aplicado": r["valor_aplicado"],
                    "rent_mes": _taxa(r["rent_mes"], 0.3), "rent_12m": _taxa(r["rent_12m"], 2.5),
                    "vencimento": r["vencimento"]})
    return out


_FLUXO_INVEST = re.compile(r"aplica|resgate|caixinha|cofrinho|\brdb\b|\bcdb\b|tesouro|invest", re.I)


def rendimento_investimentos() -> dict:
    """Quanto a carteira rendeu de verdade em cada mês, pela variação do saldo guardado a cada atualização.

    rendimento do mês = saldo no fim do mês - saldo no fim do mês anterior - (aplicações - resgates).
    As aplicações e os resgates vêm das transações do mesmo banco (item da Pluggy).
    """
    rows = store.con.execute("SELECT invest_id, item_id, data, saldo FROM invest_historico ORDER BY data").fetchall()
    if not rows:
        return {"meses": [], "inicio": None, "por_invest": {}}
    ultimo: dict[tuple[str, str], tuple[str, float]] = {}   # (mês, invest) -> (data, saldo) mais recente do mês
    item_de: dict[str, str] = {}
    for r in rows:
        m = r["data"][:7]
        ultimo[(m, r["invest_id"])] = (r["data"], r["saldo"] or 0.0)
        item_de[r["invest_id"]] = r["item_id"]
    meses = sorted({m for m, _ in ultimo})
    # aplicações (-) e resgates (+) por banco e mês
    fluxos: dict[tuple[str, str], float] = defaultdict(float)
    for t in store.con.execute(
            """SELECT t.data, t.valor, t.direcao, t.descricao, t.estabelecimento, t.categoria_origem, a.item_id
               FROM transactions t JOIN accounts a ON a.id = t.account_id
               WHERE t.data >= ? AND a.subtipo <> 'CREDIT_CARD'""", (meses[0] + "-01",)).fetchall():
        cat = store.categoria(t)
        if cat == "Investimentos" or _FLUXO_INVEST.search(_sem_acento(t["descricao"] or "")):
            fluxos[(t["data"][:7], t["item_id"])] += t["valor"] if t["direcao"] == "saida" else -t["valor"]
    out = []
    for i, m in enumerate(meses):
        saldo = sum(v[1] for (mm, _), v in ultimo.items() if mm == m)
        item = {"mes": m, "saldo": saldo, "aportes": None, "rendimento": None, "rend_pct": None, "parcial": m == hoje_br().strftime("%Y-%m")}
        if i:
            ant = meses[i - 1]
            ids = [k for (mm, k) in ultimo if mm == m and (ant, k) in ultimo]   # só quem existe nos dois meses
            if ids:
                s0 = sum(ultimo[(ant, k)][1] for k in ids)
                s1 = sum(ultimo[(m, k)][1] for k in ids)
                aportes = sum(fluxos.get((m, it), 0.0) for it in {item_de[k] for k in ids})
                rend = s1 - s0 - aportes
                item.update(aportes=aportes, rendimento=rend, rend_pct=rend / s0 if s0 else None)
        out.append(item)
    por_invest: dict[str, list[dict]] = defaultdict(list)
    for (m, k), (_, saldo) in sorted(ultimo.items()):
        por_invest[k].append({"mes": m, "saldo": saldo})
    return {"meses": out[-13:], "inicio": rows[0]["data"], "por_invest": {k: v[-13:] for k, v in por_invest.items()}}


def _sgs(serie: int, n: int) -> list[tuple[str, float]]:
    """Últimos n valores de uma série do Banco Central (SGS): [(AAAA-MM-DD, valor em %)]."""
    import urllib.request

    url = f"https://api.bcb.gov.br/dados/serie/bcdata.sgs.{serie}/dados/ultimos/{n}?formato=json"
    with urllib.request.urlopen(urllib.request.Request(url, headers={"Accept": "application/json"}), timeout=6) as r:
        dados = json.loads(r.read().decode())
    return [(f"{d['data'][6:10]}-{d['data'][3:5]}-{d['data'][0:2]}", float(d["valor"])) for d in dados]


def indicadores() -> dict | None:
    """Selic, CDI (mês a mês e em 12 meses) e IPCA do Banco Central. Guardados por 12 horas."""
    if DEMO:
        cdi = [(f"{_somar_mes_str(hoje_br(), -k)}", 0.0089 + 0.0003 * (k % 3)) for k in range(12, 0, -1)]
        return _montar_indicadores(cdi, 0.15, 0.0478, "demonstração")
    guardado = store.get_meta("indicadores")
    if guardado:
        g = json.loads(guardado)
        if datetime.fromisoformat(g["buscado_em"]) > agora_br() - timedelta(hours=12):
            return g
    try:
        cdi = [(d[:7], v / 100) for d, v in _sgs(4391, 13)]      # CDI acumulado no mês (%)
        selic = _sgs(432, 1)[-1][1] / 100                          # meta Selic (% a.a.)
        ipca12 = _sgs(13522, 1)[-1][1] / 100                       # IPCA acumulado em 12 meses (%)
    except Exception:  # noqa: BLE001  (Banco Central fora do ar: usa o último que deu certo)
        return json.loads(guardado) if guardado else None
    ind = _montar_indicadores(cdi, selic, ipca12, agora_br().isoformat())
    store.set_meta("indicadores", json.dumps(ind))
    return ind


def _somar_mes_str(d: date, n: int) -> str:
    m = d.month - 1 + n
    return f"{d.year + m // 12}-{m % 12 + 1:02d}"


def _montar_indicadores(cdi: list[tuple[str, float]], selic: float, ipca12: float, quando: str) -> dict:
    ult12 = cdi[-12:]
    acum = 1.0
    for _, v in ult12:
        acum *= 1 + v
    return {"cdi_meses": [{"mes": m, "taxa": v} for m, v in cdi], "cdi_12m": acum - 1, "selic": selic,
            "ipca_12m": ipca12, "juro_real": (1 + selic) / (1 + ipca12) - 1,
            "buscado_em": quando if quando != "demonstração" else agora_br().isoformat(), "fonte": "Banco Central (SGS)"}


def definir_orcamento(categoria: str, valor_mensal: float) -> None:
    """Cria ou altera o limite mensal de uma categoria. valor_mensal <= 0 remove o orçamento."""
    if valor_mensal <= 0:
        store.con.execute("DELETE FROM orcamentos WHERE lower(categoria)=lower(?)", (categoria,))
    else:
        store.con.execute("INSERT OR REPLACE INTO orcamentos VALUES (?,?)", (categoria.strip(), valor_mensal))
    store.con.commit()


def status_orcamentos(mes: str = "") -> list[dict]:
    """Quanto já foi gasto de cada orçamento, projeção até o fim do mês e quanto ainda dá por dia."""
    rotulo, de, ate = _mes(mes)
    cats = _por_categoria(store.transacoes(de, ate))
    hoje = hoje_br()
    atual = rotulo == hoje.strftime("%Y-%m")
    dias_mes = calendar.monthrange(int(rotulo[:4]), int(rotulo[5:]))[1]
    out = []
    for r in store.con.execute("SELECT * FROM orcamentos ORDER BY categoria"):
        gasto = next((v for c, v in cats.items() if c.lower() == r["categoria"].lower()), 0.0)
        proj = gasto / hoje.day * dias_mes if atual and hoje.day else gasto
        out.append({"categoria": r["categoria"], "limite": r["valor"], "gasto": gasto,
                    "uso": gasto / r["valor"], "projecao": proj, "atual": atual,
                    "restante_dia": (r["valor"] - gasto) / max(1, dias_mes - hoje.day + 1) if atual else None})
    return out


def criar_regra(texto: str, categoria: str) -> None:
    """Toda transação cuja descrição contenha `texto` passa a ser `categoria`, no histórico todo."""
    store.add_regra(texto, categoria)


def remover_regra(texto: str) -> int:
    return store.del_regra(texto)


def regras() -> list[dict]:
    return [{"texto": r["padrao"], "categoria": r["categoria"]}
            for r in store.con.execute("SELECT padrao, categoria FROM regras ORDER BY categoria").fetchall()]


def alertas(prev: dict | None = None) -> list[dict]:
    """Orçamentos estourando, saldo baixo, fatura vencendo, gastos fora do padrão e autorização expirando."""
    al: list[dict] = []

    def add(nivel: str, texto: str) -> None:
        al.append({"nivel": nivel, "texto": texto})

    hoje = hoje_br()
    for s in status_orcamentos():
        if s["uso"] >= 1:
            add("alto", f"Orçamento de {s['categoria']} estourado: {brl(s['gasto'])} de {brl(s['limite'])}.")
        elif s["projecao"] > s["limite"] * 1.05:
            add("medio", f"No ritmo atual, {s['categoria']} vai fechar o mês em {brl(s['projecao'])} (limite {brl(s['limite'])}).")
    for c in store.con.execute("SELECT * FROM accounts").fetchall():
        if c["tipo"] == "BANK" and c["saldo"] is not None and c["saldo"] < SALDO_MINIMO:
            add("medio", f"Saldo baixo em {c['banco']} · {c['nome']}: {brl(c['saldo'])}.")
        if c["subtipo"] == "CREDIT_CARD" and c["vencimento_fatura"]:
            d = (date.fromisoformat(c["vencimento_fatura"]) - hoje).days
            if 0 <= d <= 7:
                add("medio", f"Fatura {c['banco']} de {brl(c['saldo'])} vence em {d} dia(s) ({c['vencimento_fatura']}).")
    for it in store.con.execute("SELECT * FROM items").fetchall():
        if it["consentimento_expira"]:
            d = (date.fromisoformat(it["consentimento_expira"][:10]) - hoje).days
            if d <= 30:
                add("medio", f"Autorização do Meu Pluggy ({it['banco'] or it['id']}) expira em {d} dias.")
        if it["erro"]:
            add("alto", f"Conexão {it['banco'] or it['id']} com erro: {it['erro']}")
    # Gastos fora da curva nos últimos 10 dias (> 3x a mediana da categoria em 90 dias)
    hist = store.transacoes((hoje - timedelta(days=100)).isoformat(), (hoje - timedelta(days=11)).isoformat())
    base: dict[str, list[float]] = defaultdict(list)
    for t in hist:
        if t["direcao"] == "saida" and not t["neutra"]:
            base[t["categoria"]].append(t["valor"])
    for t in store.transacoes((hoje - timedelta(days=10)).isoformat(), hoje.isoformat()):
        b = base.get(t["categoria"])
        if t["direcao"] == "saida" and not t["neutra"] and b and len(b) >= 4 and not t.get("parcelas_total"):
            med = statistics.median(b)
            if t["valor"] > 3 * med and t["valor"] > 100:
                add("medio", f"Gasto fora do padrão: {t['descricao']} {brl(t['valor'])} em {t['data']} "
                             f"(normal em {t['categoria']}: ~{brl(med)}).")
    prev = prev or _seguro(lambda: previsao(3))
    if prev:
        vermelho = next((m for m in prev["meses"][1:4] if m["sobra"] < 0), None)
        if vermelho:
            add("medio", f"Previsão: {vermelho['mes'][5:]}/{vermelho['mes'][:4]} deve fechar no vermelho "
                         f"({brl(vermelho['sobra'])}: gastos {brl(vermelho['gastos'])}, receitas {brl(vermelho['receitas'])}).")
    ult = store.get_meta("ultima_sync")
    if not DEMO and ult and agora_br() - datetime.fromisoformat(ult) > timedelta(days=3):
        add("medio", f"Dados desatualizados desde {ult}. Clique em Atualizar.")
    return al


def _seguro(fn):
    try:
        return fn()
    except Exception:  # noqa: BLE001  (um indicador fora do ar não pode derrubar o painel)
        return None


def painel(mes: str = "") -> dict:
    """Tudo que o painel mostra, de uma vez."""
    aviso = garantir_dados()
    return {
        "aviso": aviso,
        "demo": DEMO,
        "ultima_sync": store.get_meta("ultima_sync"),
        "ultimo_resultado": json.loads(store.get_meta("ultimo_resultado") or "null"),
        "total_transacoes": store.con.execute("SELECT COUNT(*) n FROM transactions").fetchone()["n"],
        "meses": meses_disponiveis(),
        "patrimonio": patrimonio(),
        "resumo": resumo_mes(mes),
        "evolucao": evolucao_mensal(12),
        "contas": contas(),
        "conexoes": status_conexoes(),
        "orcamentos": status_orcamentos(mes),
        "recorrentes": recorrentes(),
        "parcelas": (parc := parcelas_futuras()),
        "parcelas_meses": parcelas_por_mes(parc),
        "previsao": (prev := _seguro(previsao)),
        "faturas": faturas(),
        "investimentos": investimentos(),
        "rendimento_investimentos": rendimento_investimentos(),
        "indicadores": _seguro(indicadores),
        "regras": regras(),
        "alertas": alertas(prev),
        "neutras": sorted(NEUTRAS),
    }


def gerar_planilha(de: str = "", ate: str = "", meses: int = 6):
    """Monta a planilha. Retorna (workbook, de, ate, nº de transações, nº de meses)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    hoje = hoje_br()
    if not de:
        inicio = hoje.replace(day=1)
        for _ in range(max(1, meses) - 1):
            inicio = (inicio - timedelta(days=1)).replace(day=1)
        de = inicio.isoformat()
    ate = ate or hoje.isoformat()
    txs = sorted(store.transacoes(de, ate), key=lambda t: t["data"])
    wb = Workbook()
    cab = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="1F4E78")
    fmt = '"R$" #,##0.00;[Red]-"R$" #,##0.00'

    def cabecalho(ws, nomes, larguras):
        ws.append(nomes)
        for i, w in enumerate(larguras, 1):
            ws.cell(1, i).font, ws.cell(1, i).fill = cab, fundo
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Transacoes"
    cabecalho(ws, ["Data", "Mês", "Banco", "Conta", "Descrição", "Estabelecimento", "Categoria",
                   "Tipo", "Valor", "Valor com sinal", "Conta como gasto?", "Parcela"],
              [12, 9, 12, 18, 38, 20, 24, 9, 13, 15, 16, 9])
    for t in txs:
        sinal = -t["valor"] if t["direcao"] == "saida" else t["valor"]
        ws.append([date.fromisoformat(t["data"]), t["data"][:7], t.get("banco"), t.get("conta"), t["descricao"],
                   t.get("estabelecimento"), t["categoria"], t["direcao"], t["valor"], sinal,
                   "Não" if t["neutra"] else "Sim",
                   f"{t['parcela_atual']}/{t['parcelas_total']}" if t.get("parcelas_total") else None])
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 1).number_format = "DD/MM/YYYY"
        ws.cell(r, 9).number_format = ws.cell(r, 10).number_format = fmt
    ws.auto_filter.ref = ws.dimensions

    meses_lst = sorted({t["data"][:7] for t in txs})
    matriz: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for m in meses_lst:
        for c, v in _por_categoria([t for t in txs if t["data"][:7] == m]).items():
            matriz[c][m] = v
    ws2 = wb.create_sheet("Categoria x Mes")
    cabecalho(ws2, ["Categoria"] + meses_lst + ["Média"], [26] + [13] * (len(meses_lst) + 1))
    for c in sorted(matriz, key=lambda c: -sum(matriz[c].values())):
        vals = [round(matriz[c].get(m, 0), 2) for m in meses_lst]
        ws2.append([c] + vals + [round(sum(vals) / len(vals), 2)])
    for row in ws2.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = fmt

    ws3 = wb.create_sheet("Resumo mensal")
    cabecalho(ws3, ["Mês", "Receitas", "Gastos", "Sobra", "Taxa de poupança"], [10, 14, 14, 14, 16])
    for m in meses_lst:
        rec, desp = _fluxo([t for t in txs if t["data"][:7] == m])
        ws3.append([m, rec, desp, rec - desp, (rec - desp) / rec if rec else None])
    for r in range(2, ws3.max_row + 1):
        for col in (2, 3, 4):
            ws3.cell(r, col).number_format = fmt
        ws3.cell(r, 5).number_format = "0.0%"

    return wb, de, ate, len(txs), len(meses_lst)
